"""Explainable batch resume screening CLI."""
from __future__ import annotations
import argparse, json, os, re, time
from pathlib import Path
from typing import Any
import requests

try:
    import fitz
except ImportError:
    fitz = None
try:
    from docx import Document
except ImportError:
    Document = None

WEIGHTS = {"ai_project_depth": 40, "python_backend": 30, "cloud_fullstack": 15, "github": 10, "engineering_depth": 5}
AI_TERMS = ["langchain", "langgraph", "google adk", "llamaindex", "llama index", "rag", "retrieval augmented", "vector search", "embeddings", "chromadb", "faiss", "pinecone", "agentic", "tool-calling", "tool calling", "multi-agent", "multi agent", "llm", "large language model", "openai api", "gemini api", "transformers", "hugging face", "semantic search", "prompt engineering", "generative ai", "genai"]
PYTHON_TERMS = ["python", "fastapi", "django", "flask", "pandas", "numpy", "scikit-learn", "pytorch", "tensorflow", "pytest"]
BACKEND_TERMS = ["fastapi", "django", "flask", "postgresql", "postgres", "redis", "asyncio", "async/await", "microservices", "rest api", "api development", "sqlalchemy", "celery", "kafka"]
CLOUD_TERMS = ["gcp", "google cloud", "aws", "azure", "docker", "kubernetes", "ci/cd", "github actions", "deployment", "vercel", "full-stack", "full stack", "react", "next.js"]
DEPTH_TERMS = ["unit test", "pytest", "testing", "observability", "logging", "caching", "cache", "queue", "concurrency", "retry", "retries", "failure handling", "monitoring", "architecture", "rate limit", "rollback", "latency", "optimization", "optimized", "event-driven", "event driven"]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        if fitz is None: raise RuntimeError("PyMuPDF not installed")
        with fitz.open(path) as doc:
            return "\n".join(page.get_text("text") for page in doc)
    if suffix == ".docx":
        if Document is None: raise RuntimeError("python-docx not installed")
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    if suffix == ".txt": return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}")


def unique_terms(text: str, terms: list[str]) -> list[str]:
    low = text.lower()
    return [term for term in terms if re.search(r"(?<![a-z0-9])" + re.escape(term.lower()) + r"(?![a-z0-9])", low)]


def candidate_name(text: str, filename: str) -> str:
    # Prefer a plausible first non-empty line; avoid common section headings.
    bad = {"resume", "curriculum vitae", "professional summary", "summary", "education", "technical skills", "experience", "projects"}
    for line in text.splitlines()[:12]:
        s = re.sub(r"\s+", " ", line).strip(" |•·\t")
        if not s or len(s) > 70 or "@" in s or "http" in s.lower() or re.search(r"\+?\d[\d ()-]{8,}", s): continue
        if s.lower() in bad or len(s.split()) > 5: continue
        if re.search(r"[A-Za-z]", s): return s
    return Path(filename).stem.replace("_", " ").title()


def extract_github(text: str) -> tuple[str | None, str | None]:
    urls = re.findall(r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9-]{1,39})/?", text, re.I)
    for username in urls:
        if username.lower() not in {"features", "topics", "settings", "login", "orgs", "marketplace", "sponsors"}:
            return f"https://github.com/{username}", username
    return None, None


def github_enrichment(username: str | None, cache: dict[str, Any]) -> dict[str, Any]:
    if not username:
        return {"status": "missing_profile", "summary": "No GitHub profile found in resume.", "score": 0}
    if username in cache: return cache[username]
    headers = {"Accept": "application/vnd.github+json", "User-Agent": "explainable-resume-screening-assignment"}
    token = os.getenv("GITHUB_TOKEN")
    if token: headers["Authorization"] = f"Bearer {token}"
    try:
        profile = requests.get(f"https://api.github.com/users/{username}/repos", params={"per_page": 100, "sort": "updated", "type": "owner"}, headers=headers, timeout=5)
        if profile.status_code == 404:
            result = {"status": "profile_not_found", "summary": "Public GitHub profile/repositories were not found.", "score": 0}
        elif profile.status_code == 403 or profile.status_code == 429:
            result = {"status": "rate_limited", "summary": "GitHub API rate limit or access restriction; no penalty applied beyond unavailable GitHub points.", "score": 0}
        else:
            profile.raise_for_status(); repos = profile.json()
            if not isinstance(repos, list): raise ValueError("Unexpected GitHub API response")
            cutoff = time.time() - 180 * 86400
            recent = [r for r in repos if r.get("pushed_at") and _timestamp(r["pushed_at"]) >= cutoff and not r.get("fork")]
            relevant = [r for r in repos if not r.get("fork") and ("python" in (r.get("language") or "").lower() or any(k in (r.get("name", "") + " " + r.get("description", "")).lower() for k in ["agent", "llm", "rag", "ai", "machine learning", "fastapi"]))]
            activity_pts = min(5, len(recent))
            repo_pts = min(5, len(relevant))
            score = activity_pts + repo_pts
            result = {"status": "success", "summary": f"{len(recent)} non-fork repos pushed in last 180 days; {len(relevant)} public Python/AI-relevant repos found.", "score": score, "recent_repositories": len(recent), "relevant_repositories": len(relevant)}
    except Exception as exc:
        result = {"status": "error", "summary": f"GitHub enrichment unavailable ({type(exc).__name__}); screening continued.", "score": 0}
    cache[username] = result
    return result


def _timestamp(value: str) -> float:
    from datetime import datetime
    return datetime.fromisoformat(value.replace("Z", "+00:00")).timestamp()


def score_candidate(text: str, github_score: int) -> tuple[dict[str, int], list[str], list[str], str, list[str]]:
    low = text.lower()
    ai = unique_terms(text, AI_TERMS)
    py = unique_terms(text, PYTHON_TERMS)
    backend = unique_terms(text, BACKEND_TERMS)
    cloud = unique_terms(text, CLOUD_TERMS)
    depth = unique_terms(text, DEPTH_TERMS)
    # Score based on breadth plus implementation evidence, capped to assignment weights.
    ai_score = min(40, 5 * min(len(ai), 4) + 2 * min(len([t for t in ["workflow", "orchestration", "stateful", "evaluation", "retrieval", "tool calling", "vector search", "pipeline", "production"] if t in low]), 5))
    if len(re.findall(r"\b(project|built|developed|implemented|designed|deployed|created)\b", low)) >= 3: ai_score = min(40, ai_score + 4)
    # thin-wrapper penalty when AI is only described as API integration and lacks workflow depth
    shallow = any(k in low for k in ["openai api", "gemini api", "llm api"]) and not any(k in low for k in ["rag", "retrieval", "vector", "embedding", "agent", "workflow", "evaluation", "orchestration", "tool calling"])
    if shallow: ai_score = max(0, ai_score - 10)
    python_score = min(30, 7 * min(len(py), 2) + 3 * min(len(backend), 3))
    if any(k in low for k in ["built", "developed", "implemented", "deployed"]) and py: python_score = min(30, python_score + 3)
    cloud_score = min(15, 2 * min(len(cloud), 5) + (3 if any(k in low for k in ["docker", "gcp", "aws", "azure", "deployed"]) else 0))
    depth_score = min(5, len(depth))
    breakdown = {"ai_project_depth": ai_score, "python_backend": python_score, "cloud_fullstack": cloud_score, "github": max(0, min(10, github_score)), "engineering_depth": depth_score}
    strengths=[]
    if ai: strengths.append("AI/LLM evidence: " + ", ".join(ai[:6]))
    if backend: strengths.append("Backend evidence: " + ", ".join(backend[:5]))
    if depth: strengths.append("Engineering depth: " + ", ".join(depth[:5]))
    concerns=[]
    if shallow: concerns.append("AI work may be a thin API wrapper; 10-point project-depth penalty applied")
    if len(ai) < 3: concerns.append("Limited breadth of explicit AI/agentic implementation evidence")
    if not any(x in low for x in ["postgresql", "postgres", "redis"]): concerns.append("Limited explicit PostgreSQL/Redis evidence")
    projects = []
    lines = [re.sub(r"\s+", " ", x).strip(" •-\t") for x in text.splitlines() if x.strip()]
    for i,line in enumerate(lines):
        if any(w in line.lower() for w in ["project", "built ", "developed ", "implemented ", "designed "]):
            if len(line) > 40: projects.append(line[:280])
    summary = projects[0] if projects else ("Resume mentions " + ", ".join(ai[:5]) + "." if ai else "No clear project summary extracted.")
    return breakdown, strengths, concerns, summary, ai


def screen_file(path: Path, cache: dict[str, Any], enrich_github: bool = True) -> dict[str, Any]:
    base = {"source_file": path.name, "candidate_name": path.stem.replace("_", " ").title(), "eligible": False, "rejection_reasons": [], "matched_skills": [], "evidence": [], "github_url": None, "github_enrichment": {"status": "not_checked", "summary": "Not checked for ineligible candidate.", "score": 0}}
    try:
        text = extract_text(path)
        if not text.strip(): raise ValueError("No extractable text; scanned/image-only PDF may require OCR")
        name = candidate_name(text, path.name)
        github_url, username = extract_github(text)
        py = unique_terms(text, PYTHON_TERMS)
        ai = unique_terms(text, AI_TERMS)
        # Python evidence must not be just arbitrary substring matching.
        if not py: base["rejection_reasons"].append("No evidence of Python stack")
        if not ai: base["rejection_reasons"].append("No AI/agentic project evidence")
        base.update({"candidate_name": name, "github_url": github_url, "matched_skills": sorted(set(py + ai + unique_terms(text, BACKEND_TERMS + CLOUD_TERMS))), "evidence": [x[:260] for x in [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if any(k in line.lower() for k in ai[:8] + py[:4]) and len(line.strip()) > 30][:8]]})
        if base["rejection_reasons"]:
            base.update({"eligible": False, "total_score": 0, "strengths": [], "concerns": base["rejection_reasons"], "project_summary": "Not ranked because the hard eligibility requirements were not met."})
            return base
        gh = github_enrichment(username, cache) if enrich_github else {"status": "skipped_by_option", "summary": "GitHub enrichment skipped by CLI option.", "score": 0}
        breakdown, strengths, concerns, summary, _ = score_candidate(text, gh["score"])
        total = sum(breakdown.values())
        base.update({"eligible": True, "total_score": total, "score_breakdown": breakdown, "strengths": strengths, "concerns": concerns, "project_summary": summary, "github_enrichment": gh})
        return base
    except Exception as exc:
        base.update({"parse_status": "failed", "error": f"{type(exc).__name__}: {str(exc)[:220]}", "rejection_reasons": ["Resume could not be parsed reliably"], "concerns": ["Resume could not be parsed reliably"], "project_summary": "No summary available because parsing failed.", "total_score": 0})
        return base


def run(input_dir: Path, output_file: Path, enrich_github: bool = True) -> dict[str, Any]:
    files = sorted(p for p in input_dir.rglob("*") if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".txt"})
    cache: dict[str, Any] = {}
    candidates = [screen_file(p, cache, enrich_github) for p in files]
    eligible = [c for c in candidates if c.get("eligible")]
    eligible.sort(key=lambda c: (-c["total_score"], c["candidate_name"].lower(), c["source_file"]))
    for rank, c in enumerate(eligible, 1): c["rank"] = rank
    rejected = [c for c in candidates if not c.get("eligible")]
    rejected.sort(key=lambda c: c["source_file"])
    for c in rejected: c["rank"] = None
    parsed = sum(1 for c in candidates if c.get("parse_status") != "failed")
    result = {"metadata": {"system": "AI Resume Screening & Ranking System", "scoring_weights": WEIGHTS, "scoring_method": "Deterministic evidence-based heuristic; GitHub enrichment uses public API when available", "input_directory": str(input_dir), "generated_at_epoch": int(time.time())}, "batch_summary": {"total_resumes": len(files), "successfully_parsed": parsed, "eligible": len(eligible), "rejected": len(rejected), "failed_unreadable": len(files)-parsed}, "ranked_eligible_candidates": eligible, "rejected_candidates": rejected}
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


def main():
    parser=argparse.ArgumentParser(description="Screen and rank resumes using explainable eligibility and scoring rules.")
    parser.add_argument("--input", required=True, help="Directory containing PDF, DOCX, or TXT resumes")
    parser.add_argument("--output", default="output/results.json", help="Output JSON path")
    parser.add_argument("--skip-github", action="store_true", help="Skip external GitHub API calls (useful for offline runs)")
    args=parser.parse_args()
    result=run(Path(args.input), Path(args.output), enrich_github=not args.skip_github)
    print(json.dumps(result["batch_summary"], indent=2))
    print(f"Results written to {args.output}")
if __name__ == "__main__": main()
